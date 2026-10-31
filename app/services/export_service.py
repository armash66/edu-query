"""Export service for faculty reports with anti-formula injection safeguards."""
from __future__ import annotations

import csv
import io
import re
from typing import Any

import duckdb
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from app.reports.student_breakdown import (
    StudentSubjectCounts,
    escape_like_pattern,
    get_student_breakdown,
    get_students_subject_counts_batch,
)
from app.reports.subject_failures import get_subject_failures


def _safe_float(val: Any) -> float | None:
    """Safely convert value to float or return None on conversion failure."""
    if val is None:
        return None
    if isinstance(val, (int, float)) and not isinstance(val, bool):
        return float(val)
    val_str = str(val).strip()
    if not val_str or val_str.upper() in ("N/A", "AB", "ABS", "--", "-"):
        return None
    try:
        return float(val_str)
    except (ValueError, TypeError):
        return None


def _safe_int(val: Any) -> int | None:
    """Safely convert value to int or return None on conversion failure."""
    if val is None:
        return None
    if isinstance(val, int) and not isinstance(val, bool):
        return val
    val_str = str(val).strip()
    if not val_str or val_str.upper() in ("N/A", "AB", "ABS", "--", "-"):
        return None
    try:
        return int(float(val_str))
    except (ValueError, TypeError):
        return None


def sanitize_cell_text(value: Any) -> str:
    """Sanitize cell value to prevent spreadsheet formula injection.

    If untrusted text begins with dangerous formula trigger characters ('=', '+', '-', '@', '\t', '\r'),
    prepend a single quote (') so spreadsheet engines treat it as a string literal.
    Legitimate numeric values (int, float, or pure numeric strings) are preserved without prefix.
    """
    if value is None:
        return ""
    if isinstance(value, (int, float)) and not isinstance(value, bool):
        return str(value)

    raw_str = str(value)
    if not raw_str:
        return ""

    first_char = raw_str[0]
    if first_char in ("=", "+", "-", "@", "\t", "\r"):
        if first_char in ("+", "-"):
            try:
                float(raw_str.strip())
                return raw_str
            except ValueError:
                return f"'{raw_str}"
        return f"'{raw_str}"

    return raw_str


# Retain sanitize_csv_cell for backward compatibility
sanitize_csv_cell = sanitize_cell_text


def _apply_workbook_styles(ws: openpyxl.worksheet.worksheet.Worksheet, header_row_idx: int = 1) -> None:
    """Apply professional typography, dark blue header fills, borders, frozen header, auto-filter, and auto-adjusted column widths."""
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    thin_border = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    for cell in ws[header_row_idx]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    ws.row_dimensions[header_row_idx].height = 26

    # Freeze header row so it stays pinned during scroll
    ws.freeze_panes = ws.cell(row=header_row_idx + 1, column=1)

    # Enable Excel auto-filter across header columns (must have at least one data row to be valid OOXML)
    if ws.max_column >= 1 and ws.max_row > header_row_idx:
        ws.auto_filter.ref = f"A{header_row_idx}:{get_column_letter(ws.max_column)}{ws.max_row}"

    for row in ws.iter_rows(min_row=header_row_idx + 1, max_row=ws.max_row, min_col=1, max_col=ws.max_column):
        for cell in row:
            cell.border = thin_border
            if cell.data_type != "s":
                cell.font = Font(name="Calibri", size=10)
            if isinstance(cell.value, float):
                if cell.number_format in (None, "General"):
                    cell.number_format = "0.00"

    # Compute column widths strictly from header_row_idx downwards to avoid metadata title blowup
    for col_idx in range(1, ws.max_column + 1):
        col_letter = get_column_letter(col_idx)
        max_len = 0
        for row_idx in range(header_row_idx, ws.max_row + 1):
            cell_val = ws.cell(row=row_idx, column=col_idx).value
            val_str = str(cell_val or "")
            if len(val_str) > max_len:
                max_len = len(val_str)
        ws.column_dimensions[col_letter].width = min(max(max_len + 4, 12), 45)


def format_classification_reason(reason: str | None) -> str:
    """Format technical classification reason to user-facing pass/fail status."""
    if not reason:
        return ""
    clean = str(reason).strip().upper()
    if clean in ("PASSING_LETTER_GRADE", "PASSING_GRADE_WITH_ABSENT_COMPONENT", "PASS", "PASSED"):
        return "Pass"
    if clean in ("FAILING_LETTER_GRADE", "EXPLICIT_FAIL_FLAG", "FAIL", "FAILED"):
        return "Fail"
    if clean in ("ABSENT_IN_ASSESSMENT", "ABSENT"):
        return "Absent"
    if clean == "CONFLICTING_RESULT_INDICATORS":
        return "Conflicting Indicators"
    if clean == "UNRECOGNIZED_GRADE":
        return "Unrecognized Grade"
    if clean == "MISSING_GRADE_UNRESOLVED":
        return "Missing Grade"
    return reason


def sanitize_filename(base: str) -> str:
    """Ensure filenames are deterministic, clean, and contain no filesystem paths."""
    clean = re.sub(r"[^a-zA-Z0-9_\-]+", "_", base).strip("_")
    return clean or "export"


def export_student_breakdown_csv(
    con: duckdb.DuckDBPyConnection,
    identifier: str,
) -> tuple[int, str, str | dict[str, Any]]:
    """Generate CSV export for a student's full subject performance breakdown."""
    rep = get_student_breakdown(con, identifier)
    if rep.status != "success":
        return (
            404 if rep.status == "not_found" else 409,
            "",
            rep.to_dict(),
        )

    student = rep.student
    assert student is not None

    buf = io.StringIO()
    writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)

    # 1. Report context header
    writer.writerow(["# EduQuery V3 - Student Academic Performance Breakdown"])
    writer.writerow([
        "Student ID",
        "Student Name",
        "Roll No",
        "Seat No",
        "Division",
        "Total Marks",
        "SGPI",
        "CGPA",
        "Overall Result",
    ])
    writer.writerow([
        sanitize_csv_cell(student.student_id),
        sanitize_csv_cell(student.name),
        sanitize_csv_cell(student.roll_no),
        sanitize_csv_cell(student.seat_no),
        sanitize_csv_cell(student.division),
        sanitize_csv_cell(student.total_marks),
        sanitize_csv_cell(student.sgpi),
        sanitize_csv_cell(student.cgpa),
        sanitize_csv_cell(student.overall_result),
    ])
    writer.writerow([])  # Spacer

    # 2. Detailed subject breakdown
    writer.writerow([
        "Course Code",
        "Course Name",
        "Kind",
        "Credits",
        "Component Scores",
        "Sum of Components",
        "Printed Total",
        "Total Match",
        "Letter Grade",
        "Recorded Outcome",
        "Classification Reason",
        "Flags",
        "Source Sheet",
        "Source Row",
    ])

    for s in rep.subjects:
        comp_parts = []
        for c_tag, c_detail in s.components.items():
            score_disp = str(c_detail.marks) if c_detail.marks is not None else (c_detail.flag or "N/A")
            comp_parts.append(f"{c_tag}: {score_disp}")
        comp_summary = "; ".join(comp_parts)

        writer.writerow([
            sanitize_csv_cell(s.course_code),
            sanitize_csv_cell(s.course_name),
            sanitize_csv_cell(s.kind),
            sanitize_csv_cell(s.credits),
            sanitize_csv_cell(comp_summary),
            sanitize_csv_cell(s.computed_marks_sum),
            sanitize_csv_cell(s.total_marks),
            sanitize_csv_cell("Match" if s.total_match is True else ("Mismatch" if s.total_match is False else "N/A")),
            sanitize_csv_cell(s.letter_grade),
            sanitize_csv_cell(s.recorded_result.upper()),
            sanitize_csv_cell(format_classification_reason(s.classification_reason)),
            sanitize_csv_cell(", ".join(s.flags)),
            sanitize_csv_cell(s.source_sheet),
            sanitize_csv_cell(s.source_row),
        ])

    filename = f"student_{sanitize_filename(student.student_id)}_breakdown.csv"
    return 200, filename, buf.getvalue()


def export_subject_failures_csv(
    con: duckdb.DuckDBPyConnection,
    course_code: str,
    division: str | None = None,
) -> tuple[int, str, str | dict[str, Any]]:
    """Generate CSV export for a subject's confirmed failing students roster."""
    rep = get_subject_failures(con, course_code, division=division)
    if rep.status != "success":
        return (
            404 if rep.status == "not_found" else 409,
            "",
            rep.to_dict(),
        )

    buf = io.StringIO()
    writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)

    # 1. Report context header
    writer.writerow(["# EduQuery V3 - Confirmed Subject Failures Roster"])
    writer.writerow([
        "Course Code",
        "Course Name",
        "Division Filter",
        "Total Enrolled",
        "Confirmed Failures",
        "Passing Count",
        "Unresolved Count",
        "Failure Rate (%)",
    ])
    writer.writerow([
        sanitize_csv_cell(rep.course_code),
        sanitize_csv_cell(rep.course_name),
        sanitize_csv_cell(rep.division_filter or "All Divisions"),
        sanitize_csv_cell(rep.total_enrolled),
        sanitize_csv_cell(rep.confirmed_failures_count),
        sanitize_csv_cell(rep.passed_count),
        sanitize_csv_cell(rep.unresolved_count),
        sanitize_csv_cell(round(rep.failure_rate_percent, 2)),
    ])
    writer.writerow([])  # Spacer

    # 2. Failing students roster
    writer.writerow([
        "Student ID",
        "Student Name",
        "Roll No",
        "Seat No",
        "Division",
        "Total Marks",
        "Letter Grade",
        "Flags",
        "Absent",
        "Failure Reason",
        "Component Marks",
        "Source Sheet",
        "Source Row",
    ])

    for st in rep.failing_students:
        comp_parts = []
        for c_tag, c_detail in st.components.items():
            score_disp = str(c_detail.marks) if c_detail.marks is not None else (c_detail.flag or "N/A")
            comp_parts.append(f"{c_tag}: {score_disp}")
        comp_summary = "; ".join(comp_parts)

        writer.writerow([
            sanitize_csv_cell(st.student_id),
            sanitize_csv_cell(st.name),
            sanitize_csv_cell(st.roll_no),
            sanitize_csv_cell(st.seat_no),
            sanitize_csv_cell(st.division),
            sanitize_csv_cell(st.total_marks),
            sanitize_csv_cell(st.letter_grade),
            sanitize_csv_cell(", ".join(st.flags)),
            sanitize_csv_cell("Yes" if st.is_absent else "No"),
            sanitize_csv_cell(format_classification_reason(st.failure_reason)),
            sanitize_csv_cell(comp_summary),
            sanitize_csv_cell(st.source_sheet),
            sanitize_csv_cell(st.source_row),
        ])

    div_tag = f"_{sanitize_filename(division)}" if division else ""
    filename = f"subject_{sanitize_filename(rep.course_code or 'course')}{div_tag}_failures.csv"
    return 200, filename, buf.getvalue()


def export_students_roster_csv(
    con: duckdb.DuckDBPyConnection,
    query: str | None = None,
    division: str | None = None,
    result: str | None = None,
    course_code: str | None = None,
    multi_fails: bool | None = None,
) -> tuple[str, str]:
    """Generate CSV export for full filtered student roster."""
    where_clauses: list[str] = []
    params: list[Any] = []

    if division and division.strip():
        where_clauses.append("UPPER(TRIM(division)) = UPPER(TRIM(?))")
        params.append(division.strip())

    if result and result.strip():
        r_clean = result.strip().upper()
        if "PASS" in r_clean:
            where_clauses.append("UPPER(TRIM(result)) LIKE '%PASS%'")
        elif "FAIL" in r_clean:
            where_clauses.append("UPPER(TRIM(result)) LIKE '%FAIL%'")
        elif "UNRESOLVED" in r_clean:
            where_clauses.append("NOT (UPPER(TRIM(result)) LIKE '%PASS%' OR UPPER(TRIM(result)) LIKE '%FAIL%')")
        else:
            where_clauses.append("UPPER(TRIM(result)) = UPPER(TRIM(?))")
            params.append(r_clean)

    if course_code and course_code.strip():
        where_clauses.append("student_id IN (SELECT student_id FROM enrollment WHERE UPPER(course_code) = UPPER(?))")
        params.append(course_code.strip())

    if multi_fails is True:
        where_clauses.append(
            """student_id IN (
                SELECT student_id
                FROM reconciliation
                WHERE (UPPER(TRIM(grade)) IN ('F', 'FAIL') OR UPPER(flags) LIKE '%F%')
                  AND NOT (UPPER(TRIM(grade)) IN ('O', 'A+', 'A', 'B+', 'B', 'C', 'D', 'P', 'PASS') AND UPPER(flags) LIKE '%F%')
                GROUP BY student_id
                HAVING COUNT(DISTINCT course_code) >= 2
            )"""
        )

    if query and query.strip():
        # Escape wildcards
        clean_q = query.strip().replace("!", "!!").replace("%", "!%").replace("_", "!_")
        where_clauses.append(
            "(student_id LIKE ? ESCAPE '!' OR seat_no LIKE ? ESCAPE '!' OR roll_no LIKE ? ESCAPE '!' OR name LIKE ? ESCAPE '!')"
        )
        like_term = f"%{clean_q}%"
        params.extend([like_term, like_term, like_term, like_term])

    where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""
    sql = f"""
        SELECT student_id, name, roll_no, seat_no, division, level, total_marks, result, sgpi, cgpa, source_sheet, source_row
        FROM students
        {where_sql}
        ORDER BY division, student_id
    """
    rows = con.execute(sql, params).fetchall()

    buf = io.StringIO()
    writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)

    writer.writerow(["# EduQuery V3 - Student Roster Export"])
    writer.writerow([
        "Student ID",
        "Student Name",
        "Roll No",
        "Seat No",
        "Division",
        "Level",
        "Total Marks",
        "Overall Result",
        "SGPI",
        "CGPA",
        "Source Sheet",
        "Source Row",
    ])

    for r in rows:
        writer.writerow([
            sanitize_csv_cell(r[0]),
            sanitize_csv_cell(r[1]),
            sanitize_csv_cell(r[2]),
            sanitize_csv_cell(r[3]),
            sanitize_csv_cell(r[4]),
            sanitize_csv_cell(r[5]),
            sanitize_csv_cell(r[6]),
            sanitize_csv_cell(r[7]),
            sanitize_csv_cell(r[8]),
            sanitize_csv_cell(r[9]),
            sanitize_csv_cell(r[10]),
            sanitize_csv_cell(r[11]),
        ])

    return "student_roster.csv", buf.getvalue()


def export_students_roster_excel(
    con: duckdb.DuckDBPyConnection,
    query: str | None = None,
    division: str | None = None,
    result: str | None = None,
    course_code: str | None = None,
    multi_fails: bool | None = None,
) -> tuple[str, bytes]:
    """Generate professional .xlsx export for filtered student roster."""
    where_clauses: list[str] = []
    params: list[Any] = []

    if division and division.strip():
        where_clauses.append("UPPER(TRIM(division)) = UPPER(TRIM(?))")
        params.append(division.strip())

    if result and result.strip():
        r_clean = result.strip().upper()
        if "PASS" in r_clean:
            where_clauses.append("UPPER(TRIM(result)) LIKE '%PASS%'")
        elif "FAIL" in r_clean:
            where_clauses.append("UPPER(TRIM(result)) LIKE '%FAIL%'")
        elif "UNRESOLVED" in r_clean:
            where_clauses.append("NOT (UPPER(TRIM(result)) LIKE '%PASS%' OR UPPER(TRIM(result)) LIKE '%FAIL%')")
        else:
            where_clauses.append("UPPER(TRIM(result)) = UPPER(TRIM(?))")
            params.append(r_clean)

    if course_code and course_code.strip():
        where_clauses.append("student_id IN (SELECT student_id FROM enrollment WHERE UPPER(course_code) = UPPER(?))")
        params.append(course_code.strip())

    if multi_fails is True:
        where_clauses.append(
            """student_id IN (
                SELECT student_id
                FROM reconciliation
                WHERE (UPPER(TRIM(grade)) IN ('F', 'FAIL') OR UPPER(flags) LIKE '%F%')
                  AND NOT (UPPER(TRIM(grade)) IN ('O', 'A+', 'A', 'B+', 'B', 'C', 'D', 'P', 'PASS') AND UPPER(flags) LIKE '%F%')
                GROUP BY student_id
                HAVING COUNT(DISTINCT course_code) >= 2
            )"""
        )

    if query and query.strip():
        clean_q = query.strip()
        escaped = escape_like_pattern(clean_q).lower()
        like_pattern = f"%{escaped}%"
        where_clauses.append(
            "(LOWER(name) LIKE ? ESCAPE '!' OR "
            "LOWER(student_id) LIKE ? ESCAPE '!' OR "
            "LOWER(COALESCE(seat_no, '')) LIKE ? ESCAPE '!' OR "
            "LOWER(COALESCE(roll_no, '')) LIKE ? ESCAPE '!')"
        )
        params.extend([like_pattern, like_pattern, like_pattern, like_pattern])

    where_sql = f"WHERE {' AND '.join(where_clauses)}" if where_clauses else ""
    sql = f"""
        SELECT student_id, name, roll_no, seat_no, division, level, total_marks, result, sgpi, cgpa, source_sheet, source_row
        FROM students
        {where_sql}
        ORDER BY division ASC NULLS LAST, roll_no ASC NULLS LAST, student_id ASC
    """
    rows = con.execute(sql, params).fetchall()

    student_ids = [str(r[0]) for r in rows]
    subject_counts = get_students_subject_counts_batch(con, student_ids)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Student Roster"

    headers = [
        "Student ID",
        "Student Name",
        "Roll No",
        "Seat No",
        "Division",
        "Level",
        "Total Marks",
        "Overall Result",
        "SGPI",
        "CGPA",
        "Passed Subjects",
        "Failed Subjects",
        "Absent Subjects",
        "Source Sheet",
        "Source Row",
    ]
    ws.append(headers)

    for r in rows:
        sid = str(r[0])
        counts = subject_counts.get(sid, StudentSubjectCounts(0, 0, 0, 0))
        tot_m = _safe_float(r[6])
        sgpi = _safe_float(r[8])
        cgpa = _safe_float(r[9])
        src_row = _safe_int(r[11])

        row_vals = [
            sid,
            sanitize_cell_text(r[1]),
            str(r[2]) if r[2] is not None else "",
            str(r[3]) if r[3] is not None else "",
            sanitize_cell_text(r[4] if r[4] is not None else ""),
            sanitize_cell_text(r[5] if r[5] is not None else ""),
            tot_m if tot_m is not None else "",
            sanitize_cell_text(r[7] if r[7] is not None else ""),
            sgpi if sgpi is not None else "",
            cgpa if cgpa is not None else "",
            counts.passed,
            counts.failed,
            counts.absent,
            sanitize_cell_text(r[10] if r[10] is not None else ""),
            src_row if src_row is not None else "",
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        for col_idx in (1, 3, 4):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.data_type = "s"
            cell.number_format = "@"

    _apply_workbook_styles(ws, header_row_idx=1)

    buf = io.BytesIO()
    wb.save(buf)
    is_filtered = bool(query or division or result or course_code or multi_fails)
    filename = "EduQuery_Students_Filtered.xlsx" if is_filtered else "EduQuery_Students.xlsx"
    return filename, buf.getvalue()


def export_subject_students_excel(
    con: duckdb.DuckDBPyConnection,
    course_code: str,
    division: str | None = None,
    report_type: str = "all",
) -> tuple[int, str, bytes | dict[str, Any]]:
    """Generate Excel .xlsx export for a course: all, passed, or failed students."""
    rep = get_subject_failures(con, course_identifier=course_code, division=division)
    if rep.status != "success":
        return (
            404 if rep.status == "not_found" else 409,
            "",
            rep.to_dict(),
        )

    clean_type = (report_type or "all").strip().lower()
    safe_code = sanitize_filename(rep.course_code or course_code)

    if clean_type == "passed":
        students_to_export = [s for s in rep.all_students if (s.status or "").lower() == "passed"]
        filename = f"EduQuery_{safe_code}_Passed_Students.xlsx"
        sheet_title = "Passed Students"
    elif clean_type == "failed":
        students_to_export = [s for s in rep.all_students if (s.status or "").lower() == "failed"]
        filename = f"EduQuery_{safe_code}_Failed_Students.xlsx"
        sheet_title = "Failed Students"
    elif clean_type == "absent":
        students_to_export = [s for s in rep.all_students if (s.status or "").lower() == "absent" or s.is_absent]
        filename = f"EduQuery_{safe_code}_Absent_Students.xlsx"
        sheet_title = "Absent Students"
    else:
        students_to_export = rep.all_students
        filename = f"EduQuery_{safe_code}_All_Students.xlsx"
        sheet_title = "All Students"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = sheet_title[:31]

    num_cols = 11

    # Row 1: Merged Title Banner
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=num_cols)
    cell_t = ws.cell(row=1, column=1)
    cell_t.value = f"EduQuery — Subject Marksheet Roster: {rep.course_code} - {rep.course_name}"
    cell_t.fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    cell_t.font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    cell_t.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    # Row 2: Merged Metadata
    div_label = rep.division_filter or "All Divisions"
    kind_label = (rep.kind or "Theory").title()
    credits_label = str(rep.credits or "—")
    ws.merge_cells(start_row=2, start_column=1, end_row=2, end_column=num_cols)
    cell_m = ws.cell(row=2, column=1)
    cell_m.value = f"Course: {rep.course_code} - {rep.course_name}   |   Division: {div_label}   |   Modality: {kind_label}   |   Credits: {credits_label}"
    cell_m.fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    cell_m.font = Font(name="Calibri", size=10, bold=True, color="334155")
    cell_m.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 22

    # Row 3: Merged KPI summary
    ws.merge_cells(start_row=3, start_column=1, end_row=3, end_column=num_cols)
    cell_s = ws.cell(row=3, column=1)
    cell_s.value = f"Total Enrolled: {rep.total_enrolled}   |   Passed: {rep.passed_count}   |   Failed: {rep.confirmed_failures_count}   |   Absent: {rep.absent_count}   |   Pass Rate: {rep.pass_percentage:.1f}%"
    cell_s.fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
    cell_s.font = Font(name="Calibri", size=9.5, bold=True, color="475569")
    cell_s.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[3].height = 20

    # Row 4: Spacer
    ws.append([])
    ws.row_dimensions[4].height = 8

    # Row 5: Table Headers (Keeps row 5 header index for clean styling and test compatibility)
    table_header_row = 5
    headers = [
        "Roll No",
        "Seat No",
        "Student ID",
        "Student Name",
        "Division",
        "Component Marks",
        "Total Marks",
        "Letter Grade",
        "Status",
        "Absent",
        "Remarks",
    ]
    ws.append(headers)
    ws.row_dimensions[table_header_row].height = 25

    for st in students_to_export:
        comp_parts = []
        for c_tag, c_detail in (st.components or {}).items():
            if c_tag.upper() in ('STATUS', 'RESULT') and c_detail.marks is None:
                continue
            if c_detail.marks is not None:
                m_str = f"{c_detail.marks:.0f}" if c_detail.marks.is_integer() else f"{c_detail.marks}"
                comp_parts.append(f"{c_tag}: {m_str}")
            elif c_detail.flag:
                comp_parts.append(f"{c_tag}: {c_detail.flag}")
        comp_summary = "; ".join(comp_parts) if comp_parts else "—"
        tot_m = _safe_float(st.total_marks)
        status_str = str(st.status or "").upper()
        is_abs = "Yes" if (st.is_absent or status_str == "ABSENT") else "No"
        remarks = format_classification_reason(st.classification_reason)

        row_vals = [
            str(st.roll_no or ""),
            str(st.seat_no or ""),
            str(st.student_id),
            sanitize_cell_text(st.name),
            sanitize_cell_text(st.division if st.division is not None else ""),
            sanitize_cell_text(comp_summary),
            tot_m if tot_m is not None else "",
            sanitize_cell_text(st.letter_grade if st.letter_grade is not None else ""),
            status_str,
            is_abs,
            remarks,
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        for col_idx in (1, 2, 3):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.data_type = "s"
            cell.number_format = "@"

    _apply_workbook_styles(ws, header_row_idx=table_header_row)

    # Re-apply merged banner styles
    ws.cell(row=1, column=1).font = Font(name="Calibri", size=13, bold=True, color="FFFFFF")
    ws.cell(row=1, column=1).fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    ws.cell(row=2, column=1).font = Font(name="Calibri", size=10, bold=True, color="334155")
    ws.cell(row=2, column=1).fill = PatternFill(start_color="F1F5F9", end_color="F1F5F9", fill_type="solid")
    ws.cell(row=3, column=1).font = Font(name="Calibri", size=9.5, bold=True, color="475569")
    ws.cell(row=3, column=1).fill = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

    buf = io.BytesIO()
    wb.save(buf)
    return 200, filename, buf.getvalue()


def export_subject_students_csv(
    con: duckdb.DuckDBPyConnection,
    course_code: str,
    division: str | None = None,
    report_type: str = "all",
) -> tuple[int, str, str | dict[str, Any]]:
    """Generate CSV export for a course: all, passed, failed, or absent students."""
    rep = get_subject_failures(con, course_identifier=course_code, division=division)
    if rep.status != "success":
        return (
            404 if rep.status == "not_found" else 409,
            "",
            rep.to_dict(),
        )

    clean_type = (report_type or "all").strip().lower()
    safe_code = sanitize_filename(rep.course_code or course_code)

    if clean_type == "passed":
        students_to_export = [s for s in rep.all_students if (s.status or "").lower() == "passed"]
        filename = f"EduQuery_{safe_code}_Passed_Students.csv"
    elif clean_type == "failed":
        students_to_export = [s for s in rep.all_students if (s.status or "").lower() == "failed"]
        filename = f"EduQuery_{safe_code}_Failed_Students.csv"
    elif clean_type == "absent":
        students_to_export = [s for s in rep.all_students if (s.status or "").lower() == "absent" or s.is_absent]
        filename = f"EduQuery_{safe_code}_Absent_Students.csv"
    else:
        students_to_export = rep.all_students
        filename = f"EduQuery_{safe_code}_All_Students.csv"

    buf = io.StringIO()
    writer = csv.writer(buf, quoting=csv.QUOTE_MINIMAL)

    headers = [
        "Roll No",
        "Seat No",
        "Student ID",
        "Student Name",
        "Division",
        "Component Marks",
        "Total Marks",
        "Letter Grade",
        "Status",
        "Absent",
        "Remarks",
    ]
    writer.writerow(headers)

    for st in students_to_export:
        comp_parts = []
        for c_tag, c_detail in (st.components or {}).items():
            if c_tag.upper() in ('STATUS', 'RESULT') and c_detail.marks is None:
                continue
            if c_detail.marks is not None:
                m_str = f"{c_detail.marks:.0f}" if c_detail.marks.is_integer() else f"{c_detail.marks}"
                comp_parts.append(f"{c_tag}: {m_str}")
            elif c_detail.flag:
                comp_parts.append(f"{c_tag}: {c_detail.flag}")
        comp_summary = "; ".join(comp_parts) if comp_parts else "—"
        tot_m = _safe_float(st.total_marks)
        status_str = str(st.status or "").upper()
        is_abs = "Yes" if (st.is_absent or status_str == "ABSENT") else "No"
        remarks = format_classification_reason(st.classification_reason)

        writer.writerow([
            sanitize_csv_cell(st.roll_no or ""),
            sanitize_csv_cell(st.seat_no or ""),
            sanitize_csv_cell(st.student_id),
            sanitize_cell_text(st.name),
            sanitize_cell_text(st.division if st.division is not None else ""),
            sanitize_cell_text(comp_summary),
            tot_m if tot_m is not None else "",
            sanitize_cell_text(st.letter_grade if st.letter_grade is not None else ""),
            sanitize_csv_cell(status_str),
            sanitize_csv_cell(is_abs),
            sanitize_csv_cell(remarks),
        ])

    return 200, filename, "\ufeff" + buf.getvalue()


def export_subject_students_pdf(
    con: duckdb.DuckDBPyConnection,
    course_code: str,
    division: str | None = None,
    report_type: str = "all",
) -> tuple[int, str, bytes | dict[str, Any]]:
    """Generate official printable subject evaluation marksheet roster PDF with ReportLab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas

    rep = get_subject_failures(con, course_identifier=course_code, division=division)
    if rep.status != "success":
        return (
            404 if rep.status == "not_found" else 409,
            "",
            rep.to_dict(),
        )

    clean_type = (report_type or "all").strip().lower()
    safe_code = sanitize_filename(rep.course_code or course_code)

    if clean_type == "passed":
        students_to_export = [s for s in rep.all_students if (s.status or "").lower() == "passed"]
        filename = f"EduQuery_{safe_code}_Passed_Roster.pdf"
    elif clean_type == "failed":
        students_to_export = [s for s in rep.all_students if (s.status or "").lower() == "failed"]
        filename = f"EduQuery_{safe_code}_Failed_Roster.pdf"
    elif clean_type == "absent":
        students_to_export = [s for s in rep.all_students if (s.status or "").lower() == "absent" or s.is_absent]
        filename = f"EduQuery_{safe_code}_Absent_Roster.pdf"
    else:
        students_to_export = rep.all_students
        filename = f"EduQuery_{safe_code}_All_Roster.pdf"

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=40,
        bottomMargin=42,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'RosterTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=18,
        alignment=1,
        spaceAfter=2,
        textColor=colors.HexColor("#0F172A")
    )
    sub_style = ParagraphStyle(
        'RosterSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=11,
        alignment=1,
        textColor=colors.HexColor("#475569"),
        spaceAfter=8,
    )
    th_style = ParagraphStyle('TH', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.HexColor("#FFFFFF"), alignment=1)
    th_left = ParagraphStyle('THL', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=8, leading=10, textColor=colors.HexColor("#FFFFFF"), alignment=0)

    td_style = ParagraphStyle('TD', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=9.5, textColor=colors.HexColor("#0F172A"))
    td_bold = ParagraphStyle('TDBold', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=9.5, textColor=colors.HexColor("#0F172A"))
    td_center = ParagraphStyle('TDC', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=9.5, alignment=1, textColor=colors.HexColor("#0F172A"))
    td_pass = ParagraphStyle('TDPass', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=9.5, alignment=1, textColor=colors.HexColor("#166534"))
    td_fail = ParagraphStyle('TDFail', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=9.5, alignment=1, textColor=colors.HexColor("#991B1B"))
    td_abs = ParagraphStyle('TDAbs', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=9.5, alignment=1, textColor=colors.HexColor("#C2410C"))

    story = [
        Paragraph("SUBJECT EVALUATION &amp; MARKSHEET ROSTER", title_style),
        Paragraph("Office of the Controller of Examinations &bull; Official Evaluation Record", sub_style),
        HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#0F172A"), spaceAfter=8),
    ]

    div_label = rep.division_filter or "All Divisions"
    kind_label = (rep.kind or "Theory").title()
    credits_label = str(rep.credits or "—")

    # Meta card table
    meta_rows = [
        [
            Paragraph(f"<b>Subject Code:</b> {rep.course_code}", td_style),
            Paragraph(f"<b>Subject Name:</b> {rep.course_name}", td_style),
            Paragraph(f"<b>Modality:</b> {kind_label} | <b>Credits:</b> {credits_label}", td_style),
        ],
        [
            Paragraph(f"<b>Division:</b> {div_label}", td_style),
            Paragraph(f"<b>Total Enrolled:</b> {rep.total_enrolled}  |  <b>Passed:</b> {rep.passed_count}  |  <b>Failed:</b> {rep.confirmed_failures_count}", td_style),
            Paragraph(f"<b>Pass Rate:</b> <b>{rep.pass_percentage:.1f}%</b> ({rep.passed_count}/{rep.total_enrolled})", td_style),
        ]
    ]
    t_meta = Table(meta_rows, colWidths=[150, 230, 160])
    t_meta.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor("#F8FAFC")),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#CBD5E1")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 4),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 4),
        ('LEFTPADDING', (0, 0), (-1, -1), 6),
        ('RIGHTPADDING', (0, 0), (-1, -1), 6),
    ]))
    story.append(t_meta)
    story.append(Spacer(1, 8))

    headers = [
        Paragraph("Roll No", th_style),
        Paragraph("Seat No", th_style),
        Paragraph("Student Name", th_left),
        Paragraph("Div", th_style),
        Paragraph("Component Marks", th_left),
        Paragraph("Total", th_style),
        Paragraph("Grade", th_style),
        Paragraph("Status", th_style),
    ]
    rows = [headers]

    for st in students_to_export:
        comp_parts = []
        for tag, c_detail in (st.components or {}).items():
            if tag.upper() in ('STATUS', 'RESULT') and c_detail.marks is None:
                continue
            if c_detail.marks is not None:
                m_str = f"{c_detail.marks:.0f}" if c_detail.marks.is_integer() else f"{c_detail.marks}"
                comp_parts.append(f"{tag}: {m_str}")
            elif c_detail.flag:
                comp_parts.append(f"{tag}: {c_detail.flag}")
        comp_summary = ", ".join(comp_parts) or "—"

        is_p = (st.status or '').lower() == 'passed'
        is_f = (st.status or '').lower() == 'failed'
        status_p = Paragraph("PASS", td_pass) if is_p else (Paragraph("FAIL", td_fail) if is_f else Paragraph("ABSENT", td_abs))

        tot_disp = f"{st.total_marks:.0f}" if (isinstance(st.total_marks, (int, float)) and float(st.total_marks).is_integer()) else str(st.total_marks or '—')

        rows.append([
            Paragraph(str(st.roll_no or '—'), td_center),
            Paragraph(str(st.seat_no or '—'), td_center),
            Paragraph(str(st.name or '—'), td_bold),
            Paragraph(str(st.division or '—'), td_center),
            Paragraph(comp_summary, td_style),
            Paragraph(tot_disp, td_center),
            Paragraph(str(st.letter_grade or '—'), td_center),
            status_p,
        ])

    t_roster = Table(rows, colWidths=[55, 55, 175, 30, 115, 35, 35, 40], repeatRows=1)

    style_cmds = [
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor("#1E293B")),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor("#94A3B8")),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor("#E2E8F0")),
        ('TOPPADDING', (0, 0), (-1, -1), 3),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 3),
        ('LEFTPADDING', (0, 0), (-1, -1), 4),
        ('RIGHTPADDING', (0, 0), (-1, -1), 4),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]
    for r_idx in range(1, len(rows)):
        bg = colors.HexColor("#FFFFFF") if r_idx % 2 != 0 else colors.HexColor("#F8FAFC")
        style_cmds.append(('BACKGROUND', (0, r_idx), (-1, r_idx), bg))

    t_roster.setStyle(TableStyle(style_cmds))
    story.append(t_roster)
    story.append(Spacer(1, 14))

    # Verification Signatures block
    sig_data = [
        ["____________________________", "____________________________", "____________________________"],
        ["Class Coordinator", "Subject Teacher / Faculty", "Head of Department"],
    ]
    t_sig = Table(sig_data, colWidths=[180, 180, 180])
    t_sig.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, 0), 8),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, 1), 8.5),
        ('TEXTCOLOR', (0, 1), (-1, 1), colors.HexColor("#1E293B")),
        ('PADDING', (0, 0), (-1, -1), 2),
    ]))
    story.append(KeepTogether([
        Spacer(1, 8),
        HRFlowable(width="100%", thickness=0.5, color=colors.HexColor("#CBD5E1"), spaceAfter=14),
        t_sig
    ]))

    class SubjectNumberedCanvas(canvas.Canvas):
        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_page_number(num_pages)
                canvas.Canvas.showPage(self)
            canvas.Canvas.save(self)

        def draw_page_number(self, page_count):
            self.saveState()
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#64748B"))
            if self._pageNumber > 1:
                self.drawString(36, 762, f"EduQuery Examination Cell — Subject Roster: {rep.course_code}")
                self.setStrokeColor(colors.HexColor("#CBD5E1"))
                self.setLineWidth(0.5)
                self.line(36, 756, 576, 756)

            page_str = f"Page {self._pageNumber} of {page_count}"
            self.drawString(36, 24, "EduQuery V3 · Academic Marksheet Analysis Engine · Confidential")
            self.drawRightString(576, 24, page_str)
            self.setStrokeColor(colors.HexColor("#CBD5E1"))
            self.setLineWidth(0.5)
            self.line(36, 34, 576, 34)
            self.restoreState()

    doc.build(story, canvasmaker=SubjectNumberedCanvas)
    return 200, filename, buf.getvalue()


def export_student_breakdown_excel(
    con: duckdb.DuckDBPyConnection,
    identifier: str,
) -> tuple[int, str, bytes | dict[str, Any]]:
    """Generate Excel .xlsx export for a student's full subject performance breakdown."""
    rep = get_student_breakdown(con, identifier)
    if rep.status != "success":
        return (
            404 if rep.status == "not_found" else 409,
            "",
            rep.to_dict(),
        )

    student = rep.student
    assert student is not None

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Academic Breakdown"

    ws.append([f"EduQuery — Academic Performance Breakdown: {student.name} ({student.student_id})"])
    ws.append([f"Roll No: {student.roll_no or 'N/A'} | Seat No: {student.seat_no or 'N/A'} | Division: {student.division or 'N/A'} | Result: {student.overall_result or 'N/A'}"])
    ws.append([f"Total Marks: {student.total_marks or 'N/A'} | SGPI: {student.sgpi or 'N/A'} | CGPA: {student.cgpa or 'N/A'} | Passed: {rep.passed_count} | Failed: {rep.failed_count}"])
    ws.append([])

    table_header_row = 5
    headers = [
        "Course Code",
        "Course Name",
        "Kind",
        "Credits",
        "Component Scores",
        "Computed Sum",
        "Printed Total",
        "Total Match",
        "Letter Grade",
        "Recorded Outcome",
        "Classification Reason",
        "Flags",
        "Source Sheet",
        "Source Row",
    ]
    ws.append(headers)

    for s in rep.subjects:
        comp_parts = []
        for c_tag, c_detail in s.components.items():
            score_disp = str(c_detail.marks) if c_detail.marks is not None else (c_detail.flag or "N/A")
            comp_parts.append(f"{c_tag}: {score_disp}")
        comp_summary = "; ".join(comp_parts)

        row_vals = [
            str(s.course_code),
            str(s.course_name),
            str(s.kind),
            float(s.credits) if s.credits is not None else "",
            comp_summary,
            float(s.computed_marks_sum) if s.computed_marks_sum is not None else "",
            float(s.total_marks) if s.total_marks is not None else "",
            "Match" if s.total_match is True else ("Mismatch" if s.total_match is False else "N/A"),
            str(s.letter_grade) if s.letter_grade is not None else "",
            str(s.recorded_result).upper(),
            str(format_classification_reason(s.classification_reason)),
            ", ".join(s.flags),
            str(s.source_sheet),
            int(s.source_row),
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        ws.cell(row=cur_row, column=1).data_type = "s"

    _apply_workbook_styles(ws, header_row_idx=table_header_row)

    ws.cell(row=1, column=1).font = Font(name="Calibri", size=13, bold=True, color="0F172A")
    meta_font = Font(name="Calibri", size=10, bold=True, color="334155")
    ws.cell(row=2, column=1).font = meta_font
    ws.cell(row=3, column=1).font = meta_font

    filename = f"EduQuery_Student_{sanitize_filename(student.student_id)}_Breakdown.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    return 200, filename, buf.getvalue()


def export_investigation_excel(
    con: duckdb.DuckDBPyConnection,
    where_sql: str,
    params: list[object],
    order_clause: str,
    explanation: str,
) -> tuple[str, bytes]:
    """Generate Excel .xlsx export for matching investigation cohort records."""
    select_sql = f"""
        SELECT student_id, name, roll_no, seat_no, division, level,
               total_marks, result, sgpi, cgpa, source_sheet, source_row
        FROM students
        WHERE {where_sql}
        ORDER BY {order_clause}
    """
    rows = con.execute(select_sql, params).fetchall()

    student_ids = [str(r[0]) for r in rows]
    subject_counts = get_students_subject_counts_batch(con, student_ids)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Investigation Results"

    ws.append(["EduQuery - Academic Investigation Export"])
    ws.append([f"Active Filter Query: {explanation}"])
    ws.append([f"Total Matching Students: {len(rows)}"])
    ws.append([])

    table_header_row = 5
    headers = [
        "Student ID",
        "Student Name",
        "Roll No",
        "Seat No",
        "Division",
        "Level",
        "Total Marks",
        "Overall Result",
        "SGPI",
        "CGPA",
        "Passed Subjects",
        "Failed Subjects",
        "Absent Subjects",
        "Source Sheet",
        "Source Row",
    ]
    ws.append(headers)

    for r in rows:
        sid = str(r[0])
        counts = subject_counts.get(sid, StudentSubjectCounts(0, 0, 0, 0))
        tot_m = _safe_float(r[6])
        sgpi = _safe_float(r[8])
        cgpa = _safe_float(r[9])
        src_row = _safe_int(r[11])

        row_vals = [
            sid,
            sanitize_cell_text(r[1]),
            str(r[2]) if r[2] is not None else "",
            str(r[3]) if r[3] is not None else "",
            sanitize_cell_text(r[4] if r[4] is not None else ""),
            sanitize_cell_text(r[5] if r[5] is not None else ""),
            tot_m if tot_m is not None else "",
            sanitize_cell_text(r[7] if r[7] is not None else ""),
            sgpi if sgpi is not None else "",
            cgpa if cgpa is not None else "",
            counts.passed,
            counts.failed,
            counts.absent,
            sanitize_cell_text(r[10] if r[10] is not None else ""),
            src_row if src_row is not None else "",
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        for col_idx in (1, 3, 4):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.data_type = "s"
            cell.number_format = "@"

    _apply_workbook_styles(ws, header_row_idx=table_header_row)

    ws.cell(row=1, column=1).font = Font(name="Calibri", size=13, bold=True, color="0F172A")
    meta_font = Font(name="Calibri", size=10, bold=True, color="334155")
    ws.cell(row=2, column=1).font = meta_font
    ws.cell(row=3, column=1).font = meta_font

    buf = io.BytesIO()
    wb.save(buf)
    return "EduQuery_Academic_Investigation.xlsx", buf.getvalue()


def export_absent_students_excel(
    con: duckdb.DuckDBPyConnection,
    division: str | None = None,
    course_code: str | None = None,
    component: str | None = None,
    query: str | None = None,
) -> tuple[str, bytes]:
    """Generate professional .xlsx export for verified absent student examination records."""
    where_clauses: list[str] = [
        "s.component != 'TOT1'",
        "(UPPER(TRIM(s.flag)) IN ('A', 'AB', 'ABS', 'ABSENT') OR UPPER(TRIM(s.raw_text)) IN ('A', 'AB', 'ABS', 'ABSENT'))",
    ]
    params: list[object] = []

    if division and division.strip():
        where_clauses.append("UPPER(TRIM(st.division)) = UPPER(TRIM(?))")
        params.append(division.strip())

    if course_code and course_code.strip():
        where_clauses.append("UPPER(TRIM(s.course_code)) = UPPER(TRIM(?))")
        params.append(course_code.strip())

    if component and component.strip():
        where_clauses.append("UPPER(TRIM(s.component)) = UPPER(TRIM(?))")
        params.append(component.strip())

    if query and query.strip():
        escaped = escape_like_pattern(query.strip()).lower()
        like_pattern = f"%{escaped}%"
        where_clauses.append(
            "(LOWER(st.name) LIKE ? ESCAPE '!' OR "
            "LOWER(st.student_id) LIKE ? ESCAPE '!' OR "
            "LOWER(COALESCE(st.seat_no, '')) LIKE ? ESCAPE '!' OR "
            "LOWER(COALESCE(st.roll_no, '')) LIKE ? ESCAPE '!')"
        )
        params.extend([like_pattern, like_pattern, like_pattern, like_pattern])

    where_sql = " AND ".join(where_clauses)
    select_sql = f"""
        SELECT 
            st.student_id,
            st.roll_no,
            st.seat_no,
            st.name,
            st.division,
            s.course_code,
            COALESCE(c.name, s.course_code) as course_name,
            s.component,
            s.raw_text,
            s.value,
            st.total_marks,
            st.result,
            s.source_sheet,
            s.source_row
        FROM scores s
        JOIN students st ON s.student_id = st.student_id
        LEFT JOIN courses c ON s.course_code = c.code
        WHERE {where_sql}
        ORDER BY st.division ASC NULLS LAST, st.roll_no ASC NULLS LAST, s.course_code ASC, s.component ASC
    """
    rows = con.execute(select_sql, params).fetchall()

    student_ids = list({str(r[0]) for r in rows})
    subject_counts = get_students_subject_counts_batch(con, student_ids)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Absent Students"

    filters_desc = []
    if division and division.strip():
        filters_desc.append(f"Division: {division.strip()}")
    if course_code and course_code.strip():
        filters_desc.append(f"Subject: {course_code.strip()}")
    if component and component.strip():
        filters_desc.append(f"Component: {component.strip()}")
    if query and query.strip():
        filters_desc.append(f"Query: '{query.strip()}'")
    filter_str = " | ".join(filters_desc) if filters_desc else "All Records"

    ws.append(["EduQuery - Absent Students Examination Roster"])
    ws.append([f"Filters: {filter_str} | Total Absent Records: {len(rows)} | Distinct Absent Students: {len(student_ids)}"])
    ws.append([])

    table_header_row = 4
    headers = [
        "Roll No",
        "Seat No",
        "Student ID",
        "Student Name",
        "Division",
        "Subject Code",
        "Subject Name",
        "Absent Component",
        "Absent Marker",
        "Total Marks",
        "Overall Result",
        "Passed Subjects",
        "Failed Subjects",
        "Absent Subjects",
        "Source Sheet",
        "Source Row",
    ]
    ws.append(headers)

    for r in rows:
        sid = str(r[0])
        counts = subject_counts.get(sid, StudentSubjectCounts(0, 0, 0, 0))
        tot_m = _safe_float(r[10])
        src_row = _safe_int(r[13])

        row_vals = [
            str(r[1]) if r[1] is not None else "",
            str(r[2]) if r[2] is not None else "",
            sid,
            sanitize_cell_text(r[3]),
            sanitize_cell_text(r[4] if r[4] is not None else ""),
            sanitize_cell_text(r[5]),
            sanitize_cell_text(r[6]),
            sanitize_cell_text(r[7]),
            sanitize_cell_text(r[8] if r[8] is not None else "AB"),
            tot_m if tot_m is not None else "",
            sanitize_cell_text(r[11] if r[11] is not None else ""),
            counts.passed,
            counts.failed,
            counts.absent,
            sanitize_cell_text(r[12] if r[12] is not None else ""),
            src_row if src_row is not None else "",
        ]
        ws.append(row_vals)
        cur_row = ws.max_row
        for col_idx in (1, 2, 3):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.data_type = "s"
            cell.number_format = "@"

    _apply_workbook_styles(ws, header_row_idx=table_header_row)

    ws.cell(row=1, column=1).font = Font(name="Calibri", size=13, bold=True, color="0F172A")
    meta_font = Font(name="Calibri", size=10, bold=True, color="334155")
    ws.cell(row=2, column=1).font = meta_font

    clean_subj = sanitize_filename(course_code) if course_code else ""
    filename = f"EduQuery_{clean_subj}_Absent_Students.xlsx" if clean_subj else "EduQuery_Absent_Students_Roster.xlsx"
    buf = io.BytesIO()
    wb.save(buf)
    return filename, buf.getvalue()

def export_merit_list_excel(con: duckdb.DuckDBPyConnection, anonymize: bool = False) -> tuple[str, bytes]:
    """Export complete cohort merit list ordered by CGPA, SGPI, and total marks."""
    sql = """
        SELECT 
            s.student_id, 
            s.name, 
            s.roll_no, 
            s.seat_no, 
            s.division, 
            s.level, 
            COALESCE(s.total_marks, tot.computed_total) as total_marks, 
            s.result, 
            s.sgpi, 
            s.cgpa
        FROM students s
        LEFT JOIN (
            SELECT student_id, SUM(value) as computed_total
            FROM scores
            WHERE component = 'TOT1'
            GROUP BY student_id
        ) tot ON s.student_id = tot.student_id
        WHERE UPPER(COALESCE(s.result, '')) NOT LIKE '%FAIL%'
          AND UPPER(COALESCE(s.result, '')) NOT LIKE '%ABS%'
          AND UPPER(COALESCE(s.result, '')) NOT LIKE '%BACKLOG%'
        ORDER BY 
            s.cgpa DESC NULLS LAST, 
            s.sgpi DESC NULLS LAST, 
            COALESCE(s.total_marks, tot.computed_total) DESC NULLS LAST, 
            s.seat_no ASC NULLS LAST,
            s.roll_no ASC NULLS LAST,
            s.name ASC
    """
    rows = con.execute(sql).fetchall()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Cohort Merit List"

    headers = [
        "Rank", "Student ID", "Name", "Roll No", "Seat No",
        "Division", "Level", "CGPA", "SGPI", "Total Marks", "Classification", "Result"
    ]
    ws.append(headers)

    current_rank = 1
    prev_metric = None

    for idx, r in enumerate(rows):
        cgpa_val = _safe_float(r[9])
        sgpi_val = _safe_float(r[8])
        tot_m = _safe_float(r[6])

        curr_metric = (cgpa_val, sgpi_val, tot_m)
        if prev_metric is not None and curr_metric != prev_metric:
            current_rank = idx + 1
        prev_metric = curr_metric

        metric_for_class = cgpa_val if cgpa_val is not None else sgpi_val
        classification = "Pass Class"
        if metric_for_class is not None:
            if metric_for_class >= 7.5:
                classification = "Distinction"
            elif metric_for_class >= 6.5:
                classification = "First Class"
            elif metric_for_class >= 6.0:
                classification = "Higher Second"
            elif metric_for_class >= 5.0:
                classification = "Second Class"

        name_str = f"Candidate {r[3] or r[2] or r[0]}" if anonymize else str(r[1])

        ws.append([
            current_rank,
            str(r[0]),
            sanitize_cell_text(name_str),
            str(r[2]) if r[2] is not None else "",
            str(r[3]) if r[3] is not None else "",
            sanitize_cell_text(r[4] if r[4] is not None else ""),
            sanitize_cell_text(r[5] if r[5] is not None else "UG"),
            cgpa_val if cgpa_val is not None else "",
            sgpi_val if sgpi_val is not None else "",
            tot_m if tot_m is not None else "",
            classification,
            sanitize_cell_text(r[7] if r[7] is not None else "Passed"),
        ])
        cur_row = ws.max_row
        for col_idx in (2, 4, 5):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.data_type = "s"
            cell.number_format = "@"

    _apply_workbook_styles(ws, header_row_idx=1)
    buf = io.BytesIO()
    wb.save(buf)
    return "EduQuery_Merit_List.xlsx", buf.getvalue()


def export_failed_absent_excel(con: duckdb.DuckDBPyConnection, anonymize: bool = False) -> tuple[str, bytes]:
    """Export actionable list of all failed and absent students with contact info and failing subjects."""
    sql = """
        SELECT DISTINCT s.student_id, s.name, s.roll_no, s.seat_no, s.division, s.email, s.result, s.total_marks, s.sgpi, s.cgpa
        FROM students s
        WHERE UPPER(COALESCE(s.result, '')) LIKE '%FAIL%'
           OR s.student_id IN (
               SELECT DISTINCT student_id FROM scores 
               WHERE component != 'TOT1' 
                 AND (UPPER(TRIM(flag)) IN ('A', 'AB', 'ABS', 'ABSENT')
                      OR UPPER(TRIM(raw_text)) IN ('A', 'AB', 'ABS', 'ABSENT'))
           )
        ORDER BY s.division ASC NULLS LAST, s.roll_no ASC NULLS LAST, s.student_id ASC
    """
    rows = con.execute(sql).fetchall()
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Backlog and Absent Roster"

    headers = [
        "Student ID", "Name", "Roll No", "Seat No", "Division",
        "Email", "Overall Result", "Failed Subjects Count", "Failed Courses",
        "Absent Subjects Count", "Absent Components", "SGPI", "CGPA"
    ]
    ws.append(headers)

    for r in rows:
        sid = str(r[0])
        name_str = f"Candidate {r[3] or r[2] or r[0]}" if anonymize else str(r[1])
        email_str = f"candidate_{r[3] or r[2]}@student.univ.edu" if anonymize and r[5] else str(r[5] or "")

        failed_courses_rows = con.execute("""
            SELECT course_code, grade, flags
            FROM reconciliation
            WHERE student_id = ?
              AND ((UPPER(TRIM(grade)) IN ('F', 'FAIL') OR UPPER(flags) LIKE '%F%')
                   AND NOT (UPPER(TRIM(grade)) IN ('O', 'A+', 'A', 'B+', 'B', 'C', 'D', 'P', 'PASS') AND UPPER(flags) LIKE '%F%'))
        """, [sid]).fetchall()
        failed_courses = [str(fc[0]) for fc in failed_courses_rows]

        absent_rows = con.execute("""
            SELECT course_code, component, raw_text
            FROM scores
            WHERE student_id = ? AND component != 'TOT1'
              AND (UPPER(TRIM(flag)) IN ('A', 'AB', 'ABS', 'ABSENT')
                   OR UPPER(TRIM(raw_text)) IN ('A', 'AB', 'ABS', 'ABSENT'))
        """, [sid]).fetchall()
        absent_courses = list({str(ar[0]) for ar in absent_rows})
        absent_details = ", ".join([f"{ar[0]} ({ar[1]})" for ar in absent_rows])

        ws.append([
            sid,
            sanitize_cell_text(name_str),
            str(r[2]) if r[2] is not None else "",
            str(r[3]) if r[3] is not None else "",
            sanitize_cell_text(r[4] if r[4] is not None else ""),
            sanitize_cell_text(email_str),
            sanitize_cell_text(r[6] if r[6] is not None else ""),
            len(failed_courses),
            ", ".join(failed_courses) if failed_courses else "None",
            len(absent_courses),
            absent_details if absent_details else "None",
            _safe_float(r[8]) if r[8] is not None else "",
            _safe_float(r[9]) if r[9] is not None else "",
        ])
        cur_row = ws.max_row
        for col_idx in (1, 3, 4):
            cell = ws.cell(row=cur_row, column=col_idx)
            cell.data_type = "s"
            cell.number_format = "@"

    _apply_workbook_styles(ws, header_row_idx=1)
    buf = io.BytesIO()
    wb.save(buf)
    return "EduQuery_Failed_Absent_Roster.xlsx", buf.getvalue()


def export_mail_merge_csv(con: duckdb.DuckDBPyConnection, anonymize: bool = False) -> tuple[str, bytes]:
    """Export pre-formatted RFC 4180 CSV with UTF-8 BOM for notice mail-merge."""
    sql = """
        SELECT DISTINCT s.student_id, s.name, s.roll_no, s.seat_no, s.division, s.email, s.result
        FROM students s
        WHERE UPPER(COALESCE(s.result, '')) LIKE '%FAIL%'
           OR s.student_id IN (
               SELECT DISTINCT student_id FROM scores 
               WHERE component != 'TOT1' 
                 AND (UPPER(TRIM(flag)) IN ('A', 'AB', 'ABS', 'ABSENT')
                      OR UPPER(TRIM(raw_text)) IN ('A', 'AB', 'ABS', 'ABSENT'))
           )
        ORDER BY s.division ASC NULLS LAST, s.roll_no ASC NULLS LAST, s.student_id ASC
    """
    rows = con.execute(sql).fetchall()

    buf = io.StringIO()
    writer = csv.writer(buf, lineterminator="\r\n")
    headers = [
        "Notice_Type", "Student_ID", "Student_Name", "Seat_No", "Roll_No",
        "Division", "Email", "Failed_Count", "Failed_Subjects", "Absent_Count", "Absent_Subjects", "Overall_Result"
    ]
    writer.writerow(headers)

    for r in rows:
        sid = str(r[0])
        name_str = f"Candidate {r[3] or r[2] or r[0]}" if anonymize else str(r[1])
        email_str = f"candidate_{r[3] or r[2]}@student.univ.edu" if anonymize and r[5] else str(r[5] or "")

        fc_rows = con.execute("""
            SELECT course_code FROM reconciliation
            WHERE student_id = ?
              AND ((UPPER(TRIM(grade)) IN ('F', 'FAIL') OR UPPER(flags) LIKE '%F%')
                   AND NOT (UPPER(TRIM(grade)) IN ('O', 'A+', 'A', 'B+', 'B', 'C', 'D', 'P', 'PASS') AND UPPER(flags) LIKE '%F%'))
        """, [sid]).fetchall()
        failed_courses = [str(fc[0]) for fc in fc_rows]

        ab_rows = con.execute("""
            SELECT DISTINCT course_code FROM scores
            WHERE student_id = ? AND component != 'TOT1'
              AND (UPPER(TRIM(flag)) IN ('A', 'AB', 'ABS', 'ABSENT')
                   OR UPPER(TRIM(raw_text)) IN ('A', 'AB', 'ABS', 'ABSENT'))
        """, [sid]).fetchall()
        absent_courses = [str(ab[0]) for ab in ab_rows]

        notice_type = "Backlog Notice" if failed_courses else "Absence Notice"

        writer.writerow([
            sanitize_csv_cell(notice_type),
            sanitize_csv_cell(sid),
            sanitize_csv_cell(name_str),
            sanitize_csv_cell(r[3] if r[3] is not None else ""),
            sanitize_csv_cell(r[2] if r[2] is not None else ""),
            sanitize_csv_cell(r[4] if r[4] is not None else ""),
            sanitize_csv_cell(email_str),
            len(failed_courses),
            sanitize_csv_cell(", ".join(failed_courses)),
            len(absent_courses),
            sanitize_csv_cell(", ".join(absent_courses)),
            sanitize_csv_cell(r[6] if r[6] is not None else ""),
        ])

    content_bytes = ("\ufeff" + buf.getvalue()).encode("utf-8")
    return "EduQuery_Notice_MailMerge.csv", content_bytes


def get_division_student_slips(
    con: duckdb.DuckDBPyConnection,
    division: str | None = None,
    anonymize: bool = False,
    student_id: str | None = None,
) -> list[dict[str, Any]]:
    """Retrieve complete traceable subject breakdowns for all students in a division (or entire cohort, or single student)."""
    where_clauses: list[str] = []
    params: list[Any] = []
    if student_id and student_id.strip():
        where_clauses.append("(student_id = ? OR seat_no = ?)")
        params.extend([student_id.strip(), student_id.strip()])
    elif division and division.strip().lower() not in ("all", ""):
        where_clauses.append("UPPER(TRIM(division)) = UPPER(TRIM(?))")
        params.append(division.strip())

    where_sql = ("WHERE " + " AND ".join(where_clauses)) if where_clauses else ""

    students_rows = con.execute(f"""
        SELECT student_id, name, roll_no, seat_no, division, level, total_marks, total_credits, result, sgpi, cgpa
        FROM students
        {where_sql}
        ORDER BY division ASC NULLS LAST, roll_no ASC NULLS LAST, student_id ASC
    """, params).fetchall()

    slips = []
    for s_row in students_rows:
        sid = str(s_row[0])
        name_str = f"Candidate {s_row[3] or s_row[2] or s_row[0]}" if anonymize else str(s_row[1])

        enrolled_courses = con.execute("""
            SELECT e.course_code, COALESCE(c.name, e.course_code), COALESCE(c.kind, 'Theory'), c.credits
            FROM enrollment e
            LEFT JOIN courses c ON e.course_code = c.code
            WHERE e.student_id = ?
            ORDER BY e.course_code
        """, [sid]).fetchall()

        subjects = []
        for c_code, c_name, c_kind, c_credits in enrolled_courses:
            raw_comps = con.execute("""
                SELECT component, value, flag, raw_text
                FROM scores
                WHERE student_id = ? AND course_code = ? AND component != 'TOT1'
                ORDER BY component
            """, [sid, c_code]).fetchall()

            comp_map = {}
            is_absent = False
            for comp, val, flag, raw_t in raw_comps:
                comp_map[str(comp)] = {
                    "marks": float(val) if val is not None else None,
                    "flag": str(flag) if flag is not None else None,
                    "raw_text": str(raw_t) if raw_t is not None else None,
                }
                if flag and str(flag).strip().upper() in ("A", "AB", "ABS", "ABSENT"):
                    is_absent = True

            recon = con.execute("""
                SELECT printed_total, grade
                FROM reconciliation
                WHERE student_id = ? AND course_code = ?
            """, [sid, c_code]).fetchone()

            printed_total = float(recon[0]) if recon and recon[0] is not None else None
            grade = str(recon[1]).strip() if recon and recon[1] is not None else None
            if grade and grade.upper() in ("AB", "ABS", "ABSENT"):
                is_absent = True

            status = "failed" if (grade and grade.upper() in ("F", "FAIL")) else ("absent" if is_absent else "passed")

            subjects.append({
                "course_code": str(c_code),
                "course_name": str(c_name),
                "kind": str(c_kind),
                "credits": float(c_credits) if c_credits is not None else None,
                "components": comp_map,
                "total_marks": printed_total,
                "letter_grade": grade,
                "status": status,
                "is_absent": is_absent,
            })

        slips.append({
            "student_id": sid,
            "name": name_str,
            "roll_no": str(s_row[2]) if s_row[2] is not None else "",
            "seat_no": str(s_row[3]) if s_row[3] is not None else "",
            "division": str(s_row[4] or "—"),
            "level": str(s_row[5] or "UG"),
            "total_marks": (
                _safe_float(s_row[6])
                if s_row[6] is not None
                else (sum(s["total_marks"] for s in subjects if s["total_marks"] is not None) if any(s["total_marks"] is not None for s in subjects) else None)
            ),
            "total_credits": (
                _safe_float(s_row[7])
                if s_row[7] is not None
                else (sum(s["credits"] for s in subjects if s["credits"] is not None and s["status"] == "passed") if any(s["credits"] is not None for s in subjects) else None)
            ),
            "overall_result": str(s_row[8] or "PASSES"),
            "sgpi": _safe_float(s_row[9]),
            "cgpa": _safe_float(s_row[10]),
            "subjects": subjects,
        })

    return slips


def export_analysis_dossier_pdf(con: duckdb.DuckDBPyConnection, anonymize: bool = False) -> tuple[str, bytes]:
    """Generate official 2-page institutional result analysis PDF report with ReportLab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, PageBreak, KeepTogether
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from app.services.overview_service import get_overview_summary, get_failure_summary_service

    summary = get_overview_summary(con)
    kpis = summary.kpis
    res = summary.results
    failure_rep = get_failure_summary_service(con)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=32,
        bottomMargin=32,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        alignment=1,
        spaceAfter=2,
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8.5,
        leading=12,
        alignment=1,
        textColor=colors.HexColor('#555555'),
        spaceAfter=10,
    )
    h2_style = ParagraphStyle(
        'Heading2',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=10,
        leading=13,
        spaceBefore=8,
        spaceAfter=4,
        textColor=colors.HexColor('#111111'),
    )
    th_style = ParagraphStyle('AnalysisTH', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=10)
    th_center = ParagraphStyle('AnalysisTHC', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=10, alignment=1)
    td_code = ParagraphStyle('AnalysisCode', parent=styles['Normal'], fontName='Helvetica-Bold', fontSize=7.5, leading=10)
    td_title = ParagraphStyle('AnalysisTitle', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=10)
    td_center = ParagraphStyle('AnalysisTDC', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=10, alignment=1)

    story = [
        Paragraph("INSTITUTIONAL RESULT ANALYSIS REPORT", title_style),
        Paragraph(f"Academic Examination Cell &bull; Dataset: {summary.dataset.name}", subtitle_style),
        HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#111111'), spaceAfter=10),
    ]

    # --- PAGE 1: Academic Summary, Division Distribution, Most Failed & Most Passed Subjects ---
    story.append(Paragraph("1. Academic Performance Summary", h2_style))
    kpi_data = [
        ["Total Students Enrolled", str(kpis.total_students), "Total Courses Cataloged", str(kpis.total_courses)],
        ["Passed Students", f"{kpis.passed_students} ({kpis.pass_rate_percent:.1f}%)", "Failed Students (Backlogs)", str(kpis.failed_students)],
        ["Students with Absences", str(kpis.absent_students), "Average SGPI", f"{kpis.sgpi_average:.2f}" if kpis.sgpi_average else "—"],
        ["Median SGPI", f"{kpis.sgpi_median:.2f}" if kpis.sgpi_median else "—", "Verification Flags", str(kpis.needs_attention_count)],
    ]
    t1 = Table(kpi_data, colWidths=[150, 120, 150, 120])
    t1.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#fbfbfb')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#dddddd')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#eeeeee')),
        ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('FONTNAME', (0, 0), (0, -1), 'Helvetica-Bold'),
        ('FONTNAME', (2, 0), (2, -1), 'Helvetica-Bold'),
        ('PADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t1)
    story.append(Spacer(1, 8))

    # 2. Division Breakdown Table
    story.append(Paragraph("2. Class Division Distribution", h2_style))
    div_data = [["Division", "Enrolled", "Passed", "Failed", "Pass Rate (%)"]]
    for d in res.divisions:
        failed_count = d.enrolled - d.passed
        div_data.append([d.division, str(d.enrolled), str(d.passed), str(failed_count), f"{d.pass_rate_percent:.1f}%"])
    t2 = Table(div_data, colWidths=[140, 100, 100, 100, 100])
    t2.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f0f0f0')),
        ('FONTNAME', (0, 0), (-1, 0), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 0), (-1, -1), 8),
        ('ALIGN', (1, 0), (-1, -1), 'CENTER'),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cccccc')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#eeeeee')),
        ('PADDING', (0, 0), (-1, -1), 4),
    ]))
    story.append(t2)
    story.append(Spacer(1, 8))

    # 3. Report Summary Highlights: Most Failed & Most Passed Subjects
    story.append(Paragraph("3. Key Subject Performance Extremes (Report Summary)", h2_style))
    
    # Sort subjects to isolate most failed and most passed
    all_items = list(failure_rep.summary)
    sorted_by_fail = sorted(all_items, key=lambda s: (s.confirmed_failures_count, s.failure_rate_percent), reverse=True)
    sorted_by_pass = sorted(all_items, key=lambda s: ((s.total_enrolled - s.confirmed_failures_count) / max(1, s.total_enrolled)), reverse=True)

    extremes_data = [[
        Paragraph("Performance Indicator", th_style),
        Paragraph("Code", th_style),
        Paragraph("Subject Title", th_style),
        Paragraph("Enrolled", th_center),
        Paragraph("Outcome", th_center),
        Paragraph("Rate (%)", th_center),
    ]]

    # Top 3 most failed subjects
    for rank, sub in enumerate(sorted_by_fail[:3], start=1):
        rate_str = f"{sub.failure_rate_percent:.1f}% Fail"
        extremes_data.append([
            Paragraph(f"<b>Highest Failure #{rank}</b>", td_title),
            Paragraph(sub.course_code, td_code),
            Paragraph(sub.course_name, td_title),
            Paragraph(str(sub.total_enrolled), td_center),
            Paragraph(f"{sub.confirmed_failures_count} Failed", td_center),
            Paragraph(rate_str, td_center),
        ])

    # Top 3 most passed subjects
    for rank, sub in enumerate(sorted_by_pass[:3], start=1):
        passed_cnt = sub.total_enrolled - sub.confirmed_failures_count
        pass_pct = (passed_cnt / max(1, sub.total_enrolled)) * 100.0
        extremes_data.append([
            Paragraph(f"<b>Highest Passing #{rank}</b>", td_title),
            Paragraph(sub.course_code, td_code),
            Paragraph(sub.course_name, td_title),
            Paragraph(str(sub.total_enrolled), td_center),
            Paragraph(f"{passed_cnt} Passed", td_center),
            Paragraph(f"{pass_pct:.1f}% Pass", td_center),
        ])

    t3 = Table(extremes_data, colWidths=[120, 55, 195, 50, 60, 60])
    t3.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f0f0f0')),
        ('BACKGROUND', (0, 1), (-1, 3), colors.HexColor('#fff8f8')),  # Subtle soft highlight for high failure
        ('BACKGROUND', (0, 4), (-1, -1), colors.HexColor('#f8fff8')), # Subtle soft highlight for high pass
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cccccc')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#eeeeee')),
        ('PADDING', (0, 0), (-1, -1), 3.5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t3)

    # Force clean page break so Page 2 contains ALL subjects and signatures at the very bottom
    story.append(PageBreak())

    # --- PAGE 2: Complete Subject Performance Breakdown & Signatures at bottom ---
    story.append(Paragraph("4. Complete Subject-Wise Academic Performance Analysis", h2_style))
    all_subjects_data = [[
        Paragraph("Code", th_style),
        Paragraph("Subject Title", th_style),
        Paragraph("Enrolled", th_center),
        Paragraph("Passed", th_center),
        Paragraph("Failed", th_center),
        Paragraph("Pass %", th_center),
        Paragraph("Fail %", th_center),
    ]]
    for item in failure_rep.summary:
        enrolled = item.total_enrolled
        failed = item.confirmed_failures_count
        passed = enrolled - failed
        pass_pct = (passed / enrolled * 100.0) if enrolled > 0 else 0.0
        fail_pct = item.failure_rate_percent
        all_subjects_data.append([
            Paragraph(item.course_code, td_code),
            Paragraph(item.course_name, td_title),
            Paragraph(str(enrolled), td_center),
            Paragraph(str(passed), td_center),
            Paragraph(str(failed), td_center),
            Paragraph(f"{pass_pct:.1f}%", td_center),
            Paragraph(f"{fail_pct:.1f}%", td_center),
        ])
    t_all_sub = Table(all_subjects_data, colWidths=[65, 230, 45, 45, 45, 55, 55])
    t_all_sub.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f0f0f0')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cccccc')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#eeeeee')),
        ('PADDING', (0, 0), (-1, -1), 3.5),
        ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
    ]))
    story.append(t_all_sub)
    story.append(Spacer(1, 14))

    # Key Academic Directives & Remedial Scope
    story.append(Paragraph("5. Academic Action Directives & Verification Sign-Off", h2_style))
    directives = [
        "1. Academic Remedial Classes: Mandatory mentoring sessions for all students with backlogs.",
        "2. Continuous Internal Assessment (ISA/IA): Regular attendance monitoring before End-Semester Examination.",
        "3. Official Records Reconciliation: Verified mathematically against university marksheet registry.",
    ]
    for d_text in directives:
        story.append(Paragraph(d_text, ParagraphStyle('DirText', parent=styles['Normal'], fontName='Helvetica', fontSize=7.5, leading=11, textColor=colors.HexColor('#333333'))))

    story.append(Spacer(1, 30))

    # Official Signatures (ONLY AT THE END OF PAGE 2, EXACT ORDER: Class Coordinator, Head of Department, Principal)
    sig_data = [
        ["_" * 26, "_" * 26, "_" * 26],
        ["Class Coordinator", "Head of Department", "Principal"],
    ]
    t_sig = Table(sig_data, colWidths=[180, 180, 180])
    t_sig.setStyle(TableStyle([
        ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
        ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
        ('FONTSIZE', (0, 1), (-1, 1), 8.5),
        ('PADDING', (0, 0), (-1, -1), 2),
    ]))
    story.append(KeepTogether([t_sig]))

    doc.build(story)
    return "EduQuery_Result_Analysis.pdf", buf.getvalue()


def export_student_slips_pdf(
    con: duckdb.DuckDBPyConnection,
    division: str | None = None,
    anonymize: bool = False,
    student_id: str | None = None,
) -> tuple[str, bytes]:
    """Generate multi-page printable student result slips PDF with ReportLab."""
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable, PageBreak
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors

    slips = get_division_student_slips(con, division=division, anonymize=anonymize, student_id=student_id)
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36,
    )
    styles = getSampleStyleSheet()

    h1_style = ParagraphStyle(
        'SlipH1',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=15,
        leading=19,
        alignment=1,
        spaceAfter=2,
    )
    sub_style = ParagraphStyle(
        'SlipSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=12,
        alignment=1,
        textColor=colors.HexColor('#444444'),
        spaceAfter=10,
    )
    header_style = ParagraphStyle(
        'SlipHeader',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=11,
        alignment=0,
    )
    header_center = ParagraphStyle(
        'SlipHeaderCenter',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=11,
        alignment=1,
    )
    code_cell_style = ParagraphStyle(
        'SlipCodeCell',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8,
        leading=11,
        alignment=0,
    )
    title_cell_style = ParagraphStyle(
        'SlipTitleCell',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        alignment=0,
    )
    body_center = ParagraphStyle(
        'SlipBodyCenter',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        alignment=1,
    )

    story = []
    for idx, slip in enumerate(slips):
        if idx > 0:
            story.append(PageBreak())

        story.append(Paragraph("UNIVERSITY EXAMINATION RESULT SLIP", h1_style))
        story.append(Paragraph("Academic Marksheet Statement &bull; Examination Cell", sub_style))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#000000'), spaceAfter=8))

        # Student Details Table
        meta_data = [
            [f"Student Name: {slip['name']}", f"Seat Number: {slip['seat_no'] or '—'}"],
            [f"Student ID: {slip['student_id']}", f"Roll Number: {slip['roll_no'] or '—'}"],
            [f"Division: {slip['division']}  ·  Level: {slip['level']}", f"Overall Outcome: {slip['overall_result']}"],
        ]
        t_meta = Table(meta_data, colWidths=[300, 240])
        t_meta.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('PADDING', (0, 0), (-1, -1), 2),
        ]))
        story.append(t_meta)
        story.append(Spacer(1, 8))
        story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#aaaaaa'), spaceAfter=10))

        # Courses Table with multi-line Paragraph wrapping for clean readability
        c_headers = [
            Paragraph("Code", header_style),
            Paragraph("Course Title", header_style),
            Paragraph("Credits", header_center),
            Paragraph("Marks", header_center),
            Paragraph("Grade", header_center),
            Paragraph("Status", header_center),
        ]
        c_rows = [c_headers]
        for sub in slip['subjects']:
            marks_str = str(sub['total_marks']) if sub['total_marks'] is not None else ('AB' if sub['is_absent'] else '—')
            credits_str = str(sub['credits']) if sub['credits'] is not None else '—'
            grade_str = str(sub['letter_grade']) if sub['letter_grade'] is not None else '—'
            status_str = str(sub['status']).upper()
            c_rows.append([
                Paragraph(str(sub['course_code']), code_cell_style),
                Paragraph(str(sub['course_name']), title_cell_style),
                Paragraph(credits_str, body_center),
                Paragraph(marks_str, body_center),
                Paragraph(grade_str, body_center),
                Paragraph(status_str, body_center),
            ])

        t_courses = Table(c_rows, colWidths=[70, 240, 50, 55, 55, 70])
        t_courses.setStyle(TableStyle([
            ('BACKGROUND', (0, 0), (-1, 0), colors.HexColor('#f0f0f0')),
            ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#cccccc')),
            ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#e5e5e5')),
            ('PADDING', (0, 0), (-1, -1), 5),
            ('VALIGN', (0, 0), (-1, -1), 'MIDDLE'),
        ]))
        story.append(t_courses)
        story.append(Spacer(1, 10))

        # Bottom Totals Table with real computed marks and credits
        tot_marks_str = f"{slip['total_marks']:.0f}" if slip['total_marks'] is not None else '—'
        tot_credits_str = f"{slip['total_credits']:.0f}" if slip['total_credits'] is not None else '—'
        sgpi_str = f"{slip['sgpi']:.2f}" if slip['sgpi'] is not None else '—'
        cgpa_str = f"{slip['cgpa']:.2f}" if slip['cgpa'] is not None else '—'

        totals_data = [
            [f"Total Marks: {tot_marks_str}  ·  Total Credits: {tot_credits_str}", f"SGPI: {sgpi_str}  ·  CGPA: {cgpa_str}"]
        ]
        t_tot = Table(totals_data, colWidths=[300, 240])
        t_tot.setStyle(TableStyle([
            ('FONTNAME', (0, 0), (-1, -1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 0), (-1, -1), 9),
            ('ALIGN', (1, 0), (1, -1), 'RIGHT'),
            ('PADDING', (0, 0), (-1, -1), 4),
        ]))
        story.append(t_tot)
        story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#000000'), spaceAfter=40))

        # Signature Table
        sig_data = [
            ["_" * 24, "_" * 24, "_" * 24],
            ["Controller of Examinations", "Head of Department", "Principal"],
        ]
        t_sig = Table(sig_data, colWidths=[180, 180, 180])
        t_sig.setStyle(TableStyle([
            ('ALIGN', (0, 0), (-1, -1), 'CENTER'),
            ('FONTNAME', (0, 1), (-1, 1), 'Helvetica-Bold'),
            ('FONTSIZE', (0, 1), (-1, 1), 8),
            ('PADDING', (0, 0), (-1, -1), 2),
        ]))
        story.append(t_sig)

    doc.build(story)
    if student_id:
        filename = f"EduQuery_Result_Slip_{student_id}.pdf"
    else:
        div_suffix = f"_{division}" if division else ""
        filename = f"EduQuery_Result_Slips{div_suffix}.pdf"
    return filename, buf.getvalue()




# Backward-compatible and clean alias replacing the buzzword dossier
export_analysis_report_pdf = export_analysis_dossier_pdf


def export_hod_brief_pdf(con: duckdb.DuckDBPyConnection) -> tuple[str, bytes]:
    """Generate official Department Head Executive Brief as a publication-ready PDF."""
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, HRFlowable
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from app.services.ai_service import generate_hod_brief

    brief = generate_hod_brief(con)

    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=letter,
        rightMargin=40,
        leftMargin=40,
        topMargin=40,
        bottomMargin=40,
    )
    styles = getSampleStyleSheet()

    title_style = ParagraphStyle(
        'BriefTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=16,
        leading=20,
        alignment=0,
        textColor=colors.HexColor('#111827'),
        spaceAfter=4,
    )
    subtitle_style = ParagraphStyle(
        'BriefSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        alignment=0,
        textColor=colors.HexColor('#6B7280'),
        spaceAfter=12,
    )
    section_heading = ParagraphStyle(
        'BriefSec',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=11,
        leading=15,
        textColor=colors.HexColor('#1F2937'),
        spaceBefore=14,
        spaceAfter=6,
    )
    body_style = ParagraphStyle(
        'BriefBody',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9.5,
        leading=14.5,
        textColor=colors.HexColor('#374151'),
        spaceAfter=8,
    )
    meta_label = ParagraphStyle(
        'MetaLabel',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=8.5,
        leading=11,
        textColor=colors.HexColor('#4B5563'),
    )
    meta_val = ParagraphStyle(
        'MetaVal',
        parent=styles['Normal'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=15,
        textColor=colors.HexColor('#111827'),
    )

    story = []

    # 1. Header
    story.append(Paragraph("EDUQUERY ACADEMIC RESULT MANAGEMENT", title_style))
    story.append(Paragraph("Executive Performance Brief — Department Head & Academic Review Board", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1, color=colors.HexColor('#E5E7EB'), spaceBefore=2, spaceAfter=14))

    # 2. Key Metrics Fact Grid
    facts = brief.facts_used
    total_st = facts.get('total_students', '—')
    pass_rt = f"{facts.get('pass_rate', 0.0):.1f}%" if isinstance(facts.get('pass_rate'), (int, float)) else '—'
    failed_cnt = facts.get('failed_count', '—')
    absent_cnt = facts.get('absent_count', '—')
    sgpi_avg = f"{facts.get('sgpi_average', 0.0):.2f}" if isinstance(facts.get('sgpi_average'), (int, float)) else '—'

    kpi_data = [
        [
            Paragraph("COHORT ENROLLMENT", meta_label),
            Paragraph("OVERALL PASS RATE", meta_label),
            Paragraph("BACKLOG STUDENTS", meta_label),
            Paragraph("EXAM ABSENCES", meta_label),
            Paragraph("AVERAGE SGPI", meta_label),
        ],
        [
            Paragraph(str(total_st), meta_val),
            Paragraph(str(pass_rt), meta_val),
            Paragraph(str(failed_cnt), meta_val),
            Paragraph(str(absent_cnt), meta_val),
            Paragraph(str(sgpi_avg), meta_val),
        ]
    ]

    kpi_table = Table(kpi_data, colWidths=[106, 106, 106, 106, 106])
    kpi_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F9FAFB')),
        ('BOX', (0, 0), (-1, -1), 0.75, colors.HexColor('#E5E7EB')),
        ('INNERGRID', (0, 0), (-1, -1), 0.5, colors.HexColor('#E5E7EB')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(kpi_table)
    story.append(Spacer(1, 14))

    # 3. Verified Performance Narrative
    story.append(Paragraph("EXECUTIVE PERFORMANCE ANALYSIS & RECOMMENDATIONS", section_heading))
    story.append(HRFlowable(width="100%", thickness=0.5, color=colors.HexColor('#D1D5DB'), spaceBefore=2, spaceAfter=8))

    # Split narrative by double newlines or paragraphs
    paragraphs = [p.strip() for p in brief.narrative.split('\n\n') if p.strip()]
    for p_text in paragraphs:
        # Clean markdown bold asterisks if present for clean reportlab rendering
        clean_p = p_text.replace('**', '')
        story.append(Paragraph(clean_p, body_style))

    # 4. Audit & Verification Sign-off Box
    story.append(Spacer(1, 12))
    verif_text = (
        "<b>Institutional Audit Scope:</b> All metrics and statements in this brief are deterministically verified "
        "against university marksheet tables. No extrapolated or hallucinated statistics were used."
    )
    audit_style = ParagraphStyle(
        'AuditText',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=8,
        leading=11,
        textColor=colors.HexColor('#4B5563'),
    )
    audit_table = Table([[Paragraph(verif_text, audit_style)]], colWidths=[530])
    audit_table.setStyle(TableStyle([
        ('BACKGROUND', (0, 0), (-1, -1), colors.HexColor('#F3F4F6')),
        ('BOX', (0, 0), (-1, -1), 0.5, colors.HexColor('#D1D5DB')),
        ('TOPPADDING', (0, 0), (-1, -1), 6),
        ('BOTTOMPADDING', (0, 0), (-1, -1), 6),
        ('LEFTPADDING', (0, 0), (-1, -1), 8),
        ('RIGHTPADDING', (0, 0), (-1, -1), 8),
    ]))
    story.append(audit_table)

    doc.build(story)
    pdf_bytes = buf.getvalue()
    buf.close()
    return "EduQuery_Department_Head_Brief.pdf", pdf_bytes
