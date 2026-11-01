import os
import hashlib
import json
from pathlib import Path
import openpyxl

data_dir = Path("data")
files = [
    "Result_anonymized.xlsx",
    "test1_commerce_same_layout.xlsx",
    "test1_same_layout_10courses_3divs.xlsx",
    "test2_lifesci_vague_single_row.xlsx",
    "test2_vague_single_row_batches.xlsx",
    "test3_arts_messy_sheet_per_division.xlsx",
    "test3_messy_sheet_per_division.xlsx",
]

results = []

for fname in files:
    fpath = data_dir / fname
    if not fpath.is_file():
        continue
    
    # Compute sha256
    h = hashlib.sha256()
    with open(fpath, "rb") as f:
        while chunk := f.read(65536):
            h.update(chunk)
    sha256 = h.hexdigest()
    size_bytes = fpath.stat().st_size

    # Load workbook metadata
    wb = openpyxl.load_workbook(fpath, data_only=True, read_only=True)
    sheet_names = wb.sheetnames
    
    wb_full = openpyxl.load_workbook(fpath, data_only=True)
    total_rows = sum(ws.max_row for ws in wb_full.worksheets)
    total_cols = max(ws.max_column for ws in wb_full.worksheets)
    
    results.append({
        "filename": fname,
        "sha256": sha256,
        "size_bytes": size_bytes,
        "sheets_count": len(sheet_names),
        "sheet_names": sheet_names,
        "total_rows": total_rows,
        "max_cols": total_cols,
    })

with open("paper_audit/experiments/datasets_inventory.json", "w", encoding="utf-8") as f:
    json.dump(results, f, indent=2)

print("Dataset analysis complete:")
for r in results:
    print(f"{r['filename']:<42} | Size: {r['size_bytes']:<7} | Sheets: {r['sheets_count']:<2} | SHA256: {r['sha256'][:16]}...")

