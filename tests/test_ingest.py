"""Unit tests for the EduQuery V3 ingestion engine."""
from __future__ import annotations

import tempfile
from pathlib import Path

import openpyxl
import pytest

from app.core.ingest import AmbiguousLayoutError, clean_text, detect_layout, ingest, parse_cell


def test_parse_cell_clean_numeric():
    val, flag, state = parse_cell("45")
    assert val == 45.0
    assert flag is None
    assert state == "marked"


def test_parse_cell_with_suffix_flags():
    val, flag, state = parse_cell("32F")
    assert val == 32.0
    assert flag == "F"
    assert state == "marked"

    val_e, flag_e, state_e = parse_cell("18E")
    assert val_e == 18.0
    assert flag_e == "E"
    assert state_e == "marked"


def test_parse_cell_absence_markers():
    val, flag, state = parse_cell("AB")
    assert val is None
    assert flag == "AB"
    assert state == "flagged"

    val_abs, flag_abs, state_abs = parse_cell("ABS")
    assert val_abs is None
    assert flag_abs == "ABS"
    assert state_abs == "flagged"


def test_parse_cell_blank_and_placeholders():
    val, flag, state = parse_cell("")
    assert val is None
    assert flag is None
    assert state == "not_enrolled"

    val_na, flag_na, state_na = parse_cell("N/A")
    assert val_na is None
    assert flag_na is None
    assert state_na == "uncalculated"

    val_dash, flag_dash, state_dash = parse_cell("-")
    assert val_dash is None
    assert flag_dash is None
    assert state_dash == "uncalculated"


def test_parse_cell_numeric_zero():
    val, flag, state = parse_cell("0")
    assert val == 0.0
    assert flag is None
    assert state == "marked"

    val_f, flag_f, state_f = parse_cell(0.0)
    assert val_f == 0.0
    assert flag_f is None
    assert state_f == "marked"


def test_ambiguous_layout_rejection():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "EmptySheet"
    ws.cell(1, 1, "OnlyOneCell")
    with pytest.raises(AmbiguousLayoutError):
        detect_layout(ws)


def test_missing_student_identity_rejection():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "NoIdentity"
    ws.cell(1, 1, "Subject 1")
    ws.cell(1, 2, "Subject 2")
    ws.cell(2, 1, 50)
    ws.cell(2, 2, 60)
    with pytest.raises(AmbiguousLayoutError, match="No student identity column"):
        detect_layout(ws)


def test_provenance_tracking():
    data_dir = Path(__file__).resolve().parent.parent / "data"
    sample_file = data_dir / "test1_commerce_same_layout.xlsx"
    if not sample_file.exists():
        pytest.skip(f"Sample file {sample_file} not found")

    con, rep = ingest(sample_file, db_path=":memory:")
    student_rows = con.execute("SELECT student_id, source_sheet, source_row FROM students LIMIT 5").fetchall()
    assert len(student_rows) == 5
    for sid, sheet, row in student_rows:
        assert sheet == "Sheet1"
        assert row >= 8

    score_rows = con.execute("SELECT student_id, course_code, source_sheet, source_row FROM scores LIMIT 5").fetchall()
    assert len(score_rows) == 5
    for sid, ccode, sheet, row in score_rows:
        assert sheet == "Sheet1"
        assert row >= 8


def test_embedded_course_code_and_credits_earned_summary_column(tmp_path):
    """Verify embedded course codes, credits, and summary columns like Credits Earned."""
    import openpyxl
    from app.core.ingest import ingest

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Result Sheet"
    ws.cell(1, 1, "Roll No")
    ws.cell(1, 2, "Name")
    ws.cell(1, 4, "Division")

    courses = [
        ("Strength of Materials (CE301) [4 cr]", 5),
        ("Building Materials and Construction (CE304) [3 cr]", 10),
    ]

    for title, start_col in courses:
        ws.cell(2, start_col, title)
        ws.cell(3, start_col, "INTERNAL")
        ws.cell(3, start_col+1, "EXTERNAL")
        ws.cell(3, start_col+2, "TOTAL")
        ws.cell(3, start_col+3, "GRADE")
        ws.cell(3, start_col+4, "STATUS")

    ws.cell(2, 15, "TOTAL MARKS")
    ws.cell(3, 15, "TOTAL MARKS")
    ws.cell(2, 16, "Credits Earned")
    ws.cell(3, 16, "Credits Earned")
    ws.cell(2, 17, "SGPI")
    ws.cell(3, 17, "SGPI")
    ws.cell(2, 18, "Result")
    ws.cell(3, 18, "Result")

    ws.cell(4, 1, "CV23-001")
    ws.cell(4, 2, "Qureshi, Meera")
    ws.cell(4, 4, "Batch 1")
    # Course 1: ISA=28, EXT=44, TOT=72, GR=B
    ws.cell(4, 5, 28)
    ws.cell(4, 6, 44)
    ws.cell(4, 7, 72)
    ws.cell(4, 8, "B")
    ws.cell(4, 9, "Pass")
    # Course 2: ISA=30, EXT=48, TOT=78, GR=A
    ws.cell(4, 10, 30)
    ws.cell(4, 11, 48)
    ws.cell(4, 12, 78)
    ws.cell(4, 13, "A")
    ws.cell(4, 14, "Pass")
    # Summaries
    ws.cell(4, 15, 150)
    ws.cell(4, 16, 7)
    ws.cell(4, 17, 8.5)
    ws.cell(4, 18, "Pass")

    file_path = tmp_path / "civil_test.xlsx"
    wb.save(file_path)

    con, report = ingest(str(file_path), db_path=":memory:")
    course_rows = con.execute("SELECT code, name, credits FROM courses ORDER BY code").fetchall()

    # Exactly 2 courses, NOT 3 (Credits Earned must never be a course)
    assert len(course_rows) == 2, f"Expected 2 courses, got: {course_rows}"
    codes = [c[0] for c in course_rows]
    assert "CE301" in codes
    assert "CE304" in codes
    assert "Credits Earned" not in codes

    # Credits must be 4.0 and 3.0, NOT 72.0 or 78.0
    ce304 = next(c for c in course_rows if c[0] == "CE304")
    assert ce304[1] == "Building Materials and Construction"
    assert ce304[2] == 3.0

    # Student summary columns
    st = con.execute("SELECT student_id, total_marks, total_credits, sgpi, result FROM students").fetchone()
    assert st[0] == "CV23-001"
    assert st[1] == 150.0
    assert st[2] == 7.0
    assert st[3] == 8.5
    assert st[4] == "Pass"
