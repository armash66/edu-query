"""Unit regression tests for Workbook Generalization, Ambiguous Layout Detection, and Error Handling (Phase 6 Workstream B)."""

from __future__ import annotations

from pathlib import Path
import tempfile
import openpyxl
import pytest

from app.core.ingest import AmbiguousLayoutError, ingest


def test_empty_worksheet_rejected_safely():
    """An empty workbook with no tabular data must raise AmbiguousLayoutError without crashing."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        tmp_path = Path(tf.name)

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "EmptySheet"
        wb.save(str(tmp_path))
        wb.close()

        with pytest.raises(AmbiguousLayoutError):
            ingest(tmp_path, db_path=":memory:")
    finally:
        tmp_path.unlink(missing_ok=True)


def test_ambiguous_arbitrary_table_rejected():
    """A non-academic table (e.g. library inventory) must raise AmbiguousLayoutError and never be inducted."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        tmp_path = Path(tf.name)

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Inventory"
        ws.append(["Item ID", "Description", "Location", "Quantity", "Unit Cost", "Supplier"])
        ws.append(["ITM-001", "Microscope Slides Box", "Lab 2", "50", "12.50", "Scientific Ltd"])
        ws.append(["ITM-002", "Beakers 500ml Pack", "Lab 3", "20", "24.00", "Glassworks Inc"])
        wb.save(str(tmp_path))
        wb.close()

        with pytest.raises(AmbiguousLayoutError):
            ingest(tmp_path, db_path=":memory:")
    finally:
        tmp_path.unlink(missing_ok=True)


def test_synthetic_marksheet_single_row_induction():
    """Verify single-row marksheet with proper course and assessment component headers."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        tmp_path = Path(tf.name)

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Div A"
        # Row 1: Course names above components
        ws.append([None, None, None, "CS101 - Algorithms", None, None, None, None, "CS102 - Systems", None, None, None, None])
        # Row 2: Identity + Assessment components
        ws.append(["Seat No", "Name", "Roll No", "End Sem", "Internal", "Total", "Grade", "Credits", "End Sem", "Internal", "Total", "Grade", "Credits"])
        # Row 3: Student 1
        ws.append(["SEAT01", "Alice Wonder", "101", 60, 18, 78, "A", 4.0, 55, 17, 72, "B", 4.0])
        # Row 4: Student 2
        ws.append(["SEAT02", "Bob Builder", "102", 20, 12, 32, "F", 4.0, 60, 19, 79, "A", 4.0])
        wb.save(str(tmp_path))
        wb.close()

        con, rep = ingest(tmp_path, db_path=":memory:")
        assert rep.students_count == 2
        assert rep.courses_count == 2
        assert rep.layout["rows_per_student"] == 1

        # Bob Builder should have a failure recorded in CS101
        failing_cnt = con.execute("SELECT COUNT(*) FROM reconciliation WHERE grade = 'F'").fetchone()[0]
        assert failing_cnt == 1
        con.close()
    finally:
        tmp_path.unlink(missing_ok=True)


def test_synthetic_marksheet_with_formulas():
    """Verify marksheet with formula cells (e.g. SUM) ingests properly."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        tmp_path = Path(tf.name)

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Div A"
        ws.append([None, None, None, "CS101 - Algorithms", None, None, None, None])
        ws.append(["Seat No", "Name", "Roll No", "End Sem", "Internal", "Total", "Grade", "Credits"])
        ws.append(["SEAT01", "Alice Wonder", "101", 60, 18, "=SUM(D3:E3)", "A", 4.0])
        wb.save(str(tmp_path))
        wb.close()

        con, rep = ingest(tmp_path, db_path=":memory:")
        assert rep.students_count == 1
        assert rep.courses_count == 1
        con.close()
    finally:
        tmp_path.unlink(missing_ok=True)


def test_formula_cell_without_cached_value_does_not_convert_to_zero():
    """Formula string cells without cached values must remain unparsed/null, never falsely converted to numeric zero."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        tmp_path = Path(tf.name)

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Div A"
        ws.append([None, None, None, "CS101 - Algorithms", None, None, None, None])
        ws.append(["Seat No", "Name", "Roll No", "End Sem", "Internal", "Total", "Grade", "Credits"])
        ws.append(["SEAT01", "Alice Wonder", "101", "=CALC_ESE()", 18, 70, "A", 4.0])
        wb.save(str(tmp_path))
        wb.close()

        con, rep = ingest(tmp_path, db_path=":memory:")
        score = con.execute("SELECT value, state, flag FROM scores WHERE component = 'ESE'").fetchone()
        assert score[0] is None, f"Expected value None for raw formula, got {score[0]}"
        assert score[1] in ("unparsed", "not_enrolled", "uncalculated")
        con.close()
    finally:
        tmp_path.unlink(missing_ok=True)


def test_blank_and_zero_marks_distinction():
    """Numeric zero (0.0) must be recorded as marked, while blank cells remain null/not_enrolled."""
    with tempfile.NamedTemporaryFile(suffix=".xlsx", delete=False) as tf:
        tmp_path = Path(tf.name)

    try:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Div A"
        ws.append([None, None, None, "CS101 - Algorithms", None, None, None, None])
        ws.append(["Seat No", "Name", "Roll No", "End Sem", "Internal", "Total", "Grade", "Credits"])
        # Student 1 scored numeric 0 on End Sem
        ws.append(["SEAT01", "Alice Zero", "101", 0, 15, 15, "F", 4.0])
        # Student 2 has blank cell on End Sem
        ws.append(["SEAT02", "Bob Blank", "102", None, 15, 15, "F", 4.0])
        wb.save(str(tmp_path))
        wb.close()

        con, rep = ingest(tmp_path, db_path=":memory:")
        alice_score = con.execute("SELECT value, state FROM scores WHERE student_id = '101' AND component = 'ESE'").fetchone()
        assert alice_score[0] == 0.0, "Alice numeric 0 must be 0.0"
        assert alice_score[1] == "marked", "Alice numeric 0 must have state marked"

        bob_score = con.execute("SELECT value, state FROM scores WHERE student_id = '102' AND component = 'ESE'").fetchone()
        assert bob_score[0] is None, "Bob blank cell must be None, never 0.0"
        assert bob_score[1] == "not_enrolled"
        con.close()
    finally:
        tmp_path.unlink(missing_ok=True)

